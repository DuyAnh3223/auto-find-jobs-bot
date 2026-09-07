import csv
import json
import os
import re
import tempfile
from datetime import datetime
from pathlib import Path

HEADERS = [
    "Group",
    "Group URL",
    "Từ khóa",
    "Post URL",
    "Ngày đăng",
    "Thời gian đăng gốc",
    "Ngày phát hiện",
    "Lần thấy gần nhất",
    "Nội dung",
]


def local_time(value):
    if not value:
        return ""
    return datetime.fromisoformat(value).astimezone().strftime("%d/%m/%Y %H:%M:%S %z")


def text_cell(value):
    # Untrusted post text must remain text when opened in spreadsheet software.
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def values(row):
    return [
        row["group_name"],
        row["group_url"],
        ", ".join(json.loads(row["keywords"])),
        row["url"],
        local_time(row["posted_at"]),
        row["posted_time_raw"],
        local_time(row["detected_at"]),
        local_time(row["last_seen_at"]),
        row["content"],
    ]


def export_posts(rows, destination: Path):
    """Write next to destination, then replace; a locked Excel file leaves the old file intact."""
    destination = Path(destination)
    if destination.suffix.lower() not in {".xlsx", ".csv"}:
        raise ValueError("Chỉ hỗ trợ .xlsx hoặc .csv")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=destination.parent, suffix=destination.suffix)
    os.close(fd)
    try:
        if destination.suffix.lower() == ".csv":
            with open(temporary, "w", newline="", encoding="utf-8-sig") as file:
                writer = csv.writer(file)
                writer.writerow(HEADERS)
                writer.writerows([text_cell(v) for v in values(row)] for row in rows)
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter

            book = Workbook()
            sheet = book.active
            sheet.title = "Bài đăng"
            sheet.append(HEADERS)
            for row in rows:
                cells = [text_cell(v) for v in values(row)]
                if len(cells[-1]) > 32767:
                    cells[-1] = cells[-1][:32700] + "\n[Nội dung dài: xem đầy đủ trong ứng dụng/CSV]"
                sheet.append(cells)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="175CD3")
            for index, width in enumerate([28, 40, 30, 50, 28, 24, 28, 28, 90], 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                sheet.row_dimensions[row[0].row].height = 65
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            book.save(temporary)
            book.close()
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
