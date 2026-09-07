import csv
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from monitor.core import Post
from monitor.exporter import HEADERS, export_posts
from monitor.storage import Store


def export_rows(tmp_path):
    store = Store(tmp_path / "monitor.db")
    post = Post(
        "https://www.facebook.com/groups/123/posts/456",
        '=HYPERLINK("evil")',
        "https://www.facebook.com/groups/123",
        "  =1+1\nTuyển Java, intern\x00",
        ["java intern"],
    )
    store.save_post(post)
    return store, post


def test_excel_and_csv_unicode_and_formulas(tmp_path):
    store, _ = export_rows(tmp_path)
    target = tmp_path / "results.xlsx"
    export_posts(store.posts(), target)
    book = load_workbook(target)
    sheet = book.active
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value.startswith("'=")
    assert sheet["I2"].value.startswith("'  =")
    assert "Tuyển Java" in sheet["I2"].value
    assert sheet["E2"].value is None  # unknown post time stays unknown
    book.close()
    target = tmp_path / "results.csv"
    export_posts(store.posts(), target)
    with target.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    assert len(rows) == 2
    assert rows[1][0].startswith("'=")
    assert "Tuyển Java, intern" in rows[1][-1]


def test_long_content_and_empty_export(tmp_path):
    store, post = export_rows(tmp_path)
    store.save_post(replace(post, content="a" * 40000))
    target = tmp_path / "long.xlsx"
    export_posts(store.posts(), target)
    book = load_workbook(target)
    assert "xem đầy đủ" in book.active["I2"].value
    book.close()
    export_posts(store.posts(), tmp_path / "long.csv")
    assert "a" * 40000 in (tmp_path / "long.csv").read_text(encoding="utf-8-sig")
    export_posts([], target)
    book = load_workbook(target)
    assert book.active.max_row == 1
    book.close()


def test_failed_replace_preserves_previous_export(tmp_path, monkeypatch):
    target = tmp_path / "results.xlsx"
    target.write_bytes(b"previous file")

    def locked(*args):
        raise PermissionError("file is open in Excel")

    monkeypatch.setattr("monitor.exporter.os.replace", locked)
    with pytest.raises(PermissionError):
        export_posts([], target)
    assert target.read_bytes() == b"previous file"
    assert list(tmp_path.iterdir()) == [target]
